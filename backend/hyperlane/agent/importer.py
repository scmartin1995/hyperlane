"""Excel / CSV import: heuristic column mapping (refined by Claude when available), then a bulk create
that rebuilds epic -> feature -> story hierarchies from parent/epic/feature columns."""

from __future__ import annotations

import csv
import difflib
import io
import logging
import re
import secrets
from datetime import date, datetime
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..config import ROOT
from ..models import Project, User, WorkItem
from ..services.events import bus
from . import _compat as c
from .prompts import IMPORT_SYSTEM
from .provider import ProviderError, get_provider

log = logging.getLogger(__name__)

UPLOAD_DIR = ROOT / "data" / "uploads"
PREVIEW_ROWS = 20
FIELDS = [
    "title",
    "description",
    "type",
    "status",
    "priority",
    "points",
    "assignee",
    "parent",
    "epic",
    "feature",
    "labels",
    "due_date",
    "acceptance_criteria",
    "key",
]
ALIASES: dict[str, list[str]] = {
    "title": ["title", "summary", "name", "story", "item", "subject", "task", "work item", "user story"],
    "description": ["description", "details", "desc", "notes", "body"],
    "type": ["type", "issue type", "issuetype", "kind", "category", "work item type"],
    "status": ["status", "state", "column", "stage", "workflow"],
    "priority": ["priority", "prio", "severity", "importance"],
    "points": ["points", "story points", "estimate", "size", "sp", "effort", "story point estimate", "estimation"],
    "assignee": ["assignee", "owner", "assigned to", "developer", "responsible", "who"],
    "parent": ["parent", "parent key", "parent id", "parent summary", "belongs to"],
    "epic": ["epic", "epic name", "epic link", "theme", "initiative", "objective"],
    "feature": ["feature", "feature name", "capability", "component"],
    "labels": ["labels", "tags", "label", "tag", "keywords"],
    "due_date": ["due date", "due", "deadline", "target date", "due_date", "end date"],
    "acceptance_criteria": ["acceptance criteria", "acceptance", "ac", "definition of done", "criteria"],
    "key": ["key", "id", "issue key", "ticket", "ref", "reference"],
}


class ColumnMapping(BaseModel):
    mapping: dict[str, str | None] = Field(default_factory=dict, description="field -> column header or null")
    notes: list[str] = Field(default_factory=list)


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


def _norm(s: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").strip().lower()).strip()


def _cell(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def load_table(file_bytes: bytes, filename: str, sheet: str | None = None) -> tuple[list[str], str, list[str], list[dict]]:
    """Return (sheet names, chosen sheet, columns, rows as dicts)."""
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".txt", ".tsv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        reader = csv.reader(io.StringIO(text), dialect=dialect)
        raw_rows = [r for r in reader]
        return ["Sheet1"], "Sheet1", *_rows_from_matrix(raw_rows)
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = wb.sheetnames
    chosen = sheet if sheet in sheets else sheets[0]
    ws = wb[chosen]
    raw_rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return sheets, chosen, *_rows_from_matrix(raw_rows)


def _rows_from_matrix(raw_rows: list[list[Any]]) -> tuple[list[str], list[dict]]:
    raw_rows = [r for r in raw_rows if any(v not in (None, "") for v in r)]
    if not raw_rows:
        return [], []
    # header = first row with >= 2 non-empty string cells
    header_idx = 0
    for i, r in enumerate(raw_rows[:10]):
        if sum(1 for v in r if isinstance(v, str) and v.strip()) >= 2:
            header_idx = i
            break
    header = raw_rows[header_idx]
    columns: list[str] = []
    for i, h in enumerate(header):
        name = str(h).strip() if h not in (None, "") else f"Column {i + 1}"
        while name in columns:
            name += "_"
        columns.append(name)
    rows: list[dict] = []
    for r in raw_rows[header_idx + 1 :]:
        row = {columns[i]: (_cell(r[i]) if i < len(r) else None) for i in range(len(columns))}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return columns, rows


# ---------------------------------------------------------------------------
# Mapping
# ---------------------------------------------------------------------------


def heuristic_mapping(columns: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {f: None for f in FIELDS}
    used: set[str] = set()
    normed = {col: _norm(col) for col in columns}
    # exact alias matches first (in field priority order), then fuzzy
    for field in FIELDS:
        for col in columns:
            if col in used:
                continue
            if normed[col] in ALIASES[field]:
                mapping[field] = col
                used.add(col)
                break
    for field in FIELDS:
        if mapping[field]:
            continue
        best, best_score = None, 0.0
        for col in columns:
            if col in used:
                continue
            for alias in ALIASES[field]:
                score = difflib.SequenceMatcher(None, normed[col], alias).ratio()
                if normed[col].startswith(alias) or alias in normed[col].split():
                    score = max(score, 0.85)
                if score > best_score:
                    best, best_score = col, score
        if best and best_score >= 0.8:
            mapping[field] = best
            used.add(best)
    return mapping


async def refine_mapping(
    columns: list[str], sample_rows: list[dict], mapping: dict[str, str | None]
) -> tuple[dict[str, str | None], list[str]]:
    provider = get_provider()
    if not provider.available:
        return mapping, []
    prompt = (
        f"Columns: {columns}\n"
        f"Sample rows (first {len(sample_rows)}): {sample_rows}\n"
        f"Heuristic mapping so far: {mapping}\n"
        f"Fields to map: {FIELDS}. 'epic'/'feature' are names of parent items to create; 'parent' is a key or title "
        "of an existing parent; 'key' is an external id we keep as a reference only.\n"
        "Return the full corrected mapping (field -> exact column header or null) and short notes about anything odd."
    )
    try:
        data = await provider.complete(IMPORT_SYSTEM, prompt, json_schema=ColumnMapping)
    except ProviderError as exc:
        return mapping, [f"Agent mapping skipped: {exc}"]
    assert isinstance(data, dict)
    refined = dict(mapping)
    for field, col in (data.get("mapping") or {}).items():
        if field in FIELDS and (col is None or col in columns):
            refined[field] = col
    return refined, list(data.get("notes") or [])


# ---------------------------------------------------------------------------
# Preview / commit
# ---------------------------------------------------------------------------


def _save_upload(file_bytes: bytes, filename: str) -> str:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    ext = Path(filename).suffix.lower() or ".xlsx"
    if ext not in (".xlsx", ".xlsm", ".csv", ".tsv", ".txt"):
        raise ValueError("Only .xlsx, .xlsm, .csv or .tsv files can be imported")
    token = secrets.token_hex(8)
    (UPLOAD_DIR / f"{token}{ext}").write_bytes(file_bytes)
    return token


def _find_upload(token: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{16}", token or ""):
        raise ValueError("Bad file token")
    for p in UPLOAD_DIR.glob(f"{token}.*"):
        return p
    raise ValueError("Uploaded file not found; upload it again")


async def preview(project: Project, file_bytes: bytes, filename: str, sheet: str | None = None) -> dict[str, Any]:
    token = _save_upload(file_bytes, filename)
    sheets, chosen, columns, rows = load_table(file_bytes, filename, sheet)
    notes: list[str] = []
    if not columns:
        notes.append("No header row found.")
    mapping = heuristic_mapping(columns)
    if columns:
        mapping, agent_notes = await refine_mapping(columns, rows[:5], mapping)
        notes.extend(agent_notes)
    if not mapping.get("title"):
        notes.append("Couldn't find a title column; pick one before importing.")
    return {
        "file_token": token,
        "sheet": chosen,
        "sheets": sheets,
        "columns": columns,
        "rows": rows[:PREVIEW_ROWS],
        "total_rows": len(rows),
        "mapping": mapping,
        "notes": notes,
    }


def _fuzzy_key(value: Any, options: list[dict], *, default: str, extra: dict[str, str] | None = None) -> str:
    """Map a free-text value onto one of the project's keys (by key, name, plural; then fuzzy)."""
    v = _norm(value)
    if not v:
        return default
    if extra and v in extra:
        return extra[v]
    for opt in options:
        if v in (_norm(opt.get("key")), _norm(opt.get("name")), _norm(opt.get("plural"))):
            return opt["key"]
    names = {_norm(o.get("name")): o["key"] for o in options} | {_norm(o.get("key")): o["key"] for o in options}
    close = difflib.get_close_matches(v, list(names), n=1, cutoff=0.75)
    return names[close[0]] if close else default


_PRIORITY_ALIASES = {
    "highest": "critical", "blocker": "critical", "urgent": "critical", "p0": "critical", "p1": "critical", "1": "critical",
    "high": "high", "p2": "high", "2": "high", "major": "high",
    "medium": "medium", "normal": "medium", "p3": "medium", "3": "medium", "moderate": "medium",
    "low": "low", "minor": "low", "p4": "low", "4": "low", "lowest": "low", "trivial": "low", "p5": "low", "5": "low",
    "none": "none", "": "none",
}
_PRIORITY_OPTS = [{"key": k, "name": k} for k in ("critical", "high", "medium", "low", "none")]


def _parse_points(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        m = re.search(r"\d+(?:\.\d+)?", str(v))
        return float(m.group()) if m else None


def _parse_date(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d", "%b %d, %Y", "%d %b %Y"):
        try:
            return datetime.strptime(s[: len(fmt) + 4] if "T" in fmt else s, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        return None


def _split_labels(v: Any) -> list[str]:
    if v in (None, ""):
        return []
    return [p.strip() for p in re.split(r"[,;|/]+", str(v)) if p.strip()][:10]


def _handle_from(name: str, db: Session) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "user"
    handle, n = base, 2
    while db.scalar(select(User).where(User.handle == handle)) is not None:
        handle = f"{base}-{n}"
        n += 1
    return handle


def commit(
    db: Session,
    project: Project,
    file_token: str,
    sheet: str | None,
    mapping: dict[str, str | None],
    create_missing_users: bool,
    user: User | None,
) -> dict[str, Any]:
    path = _find_upload(file_token)
    _, chosen, columns, rows = load_table(path.read_bytes(), path.name, sheet)
    warnings: list[str] = []
    mapping = {f: (col if col in columns else None) for f, col in (mapping or {}).items() if f in FIELDS}
    if not mapping.get("title"):
        raise ValueError("Mapping needs a 'title' column")

    types = project.item_types or []
    workflow = project.workflow or []
    leaf = c.leaf_type(project)
    default_status = workflow[0]["key"] if workflow else "backlog"
    status_extra: dict[str, str] = {}
    for s in workflow:  # "done"/"in progress"/"todo" category words map to the first state of that category
        status_extra.setdefault(_norm(s.get("category")), s["key"])
        status_extra.setdefault(_norm(s.get("category")).replace(" ", ""), s["key"])
    status_extra.setdefault("to do", status_extra.get("todo", default_status))
    status_extra.setdefault("open", default_status)
    status_extra.setdefault("closed", status_extra.get("done", default_status))
    epic_type = c.type_for_level(project, 0)
    feature_type = c.type_for_level(project, 1) or epic_type
    level_of = {t["key"]: int(t.get("level", 0)) for t in types}

    def get(row: dict, field: str) -> Any:
        col = mapping.get(field)
        return row.get(col) if col else None

    # existing items by title (for epic/feature dedupe and parent lookups)
    existing = {
        _norm(it.title): it
        for it in db.scalars(
            select(WorkItem).where(WorkItem.project_id == project.id, WorkItem.is_archived.is_(False))
        ).all()
    }
    containers: dict[str, dict[str, Any]] = {}  # temp_id -> spec, insertion ordered (parents first)
    specs: list[dict[str, Any]] = []
    users_cache: dict[str, str | None] = {}
    skipped = 0

    def ensure_container(title: Any, type_key: str | None, parent_temp: str | None, parent_id: str | None):
        """Return (temp_id, existing_id) for an epic/feature named in a row."""
        t = str(title).strip()
        if not t or not type_key:
            return None, None
        found = existing.get(_norm(t))
        if found is not None:
            return None, found.id
        temp = f"{type_key}:{_norm(t)}"
        if temp not in containers:
            spec: dict[str, Any] = {"temp_id": temp, "type": type_key, "title": t[:300], "source": "excel"}
            if parent_temp:
                spec["parent_temp_id"] = parent_temp
            elif parent_id:
                spec["parent_id"] = parent_id
            containers[temp] = spec
        return temp, None

    def resolve_assignee(value: Any) -> str | None:
        v = str(value or "").strip()
        if not v:
            return None
        if v in users_cache:
            return users_cache[v]
        u = c.resolve_user(db, v)
        if u is None:
            close = difflib.get_close_matches(
                v.lower(), [x.name.lower() for x in db.scalars(select(User)).all()], n=1, cutoff=0.85
            )
            if close:
                u = db.scalar(select(User).where(func.lower(User.name) == close[0]))
        if u is None and create_missing_users:
            u = User(name=v[:120], handle=_handle_from(v, db), role="developer")
            db.add(u)
            db.flush()
            bus.emit("user.created", entity_type="user", entity_id=u.id, data={"name": u.name, "source": "excel"})
            warnings.append(f"Created user '{v}'")
        elif u is None:
            if f"Unknown assignee '{v}' left unassigned" not in warnings:
                warnings.append(f"Unknown assignee '{v}' left unassigned")
        users_cache[v] = u.id if u else None
        return users_cache[v]

    for idx, row in enumerate(rows, start=2):
        title = get(row, "title")
        if title in (None, ""):
            skipped += 1
            continue
        epic_temp, epic_id = ensure_container(get(row, "epic"), epic_type, None, None)
        feature_temp, feature_id = (None, None)
        if get(row, "feature") not in (None, ""):
            feature_temp, feature_id = ensure_container(get(row, "feature"), feature_type, epic_temp, epic_id)
        parent_temp: str | None = None
        parent_id: str | None = None
        pval = get(row, "parent")
        if pval not in (None, ""):
            found = c.resolve_item(db, str(pval)) or existing.get(_norm(pval))
            if found is not None and found.project_id == project.id:
                parent_id = found.id
            else:
                # a parent named in this same file (by its title or external key)
                match = next(
                    (s for s in specs if _norm(s["title"]) == _norm(pval) or s.get("_ext_key") == _norm(pval)), None
                )
                if match is not None:
                    parent_temp = match["temp_id"]
                else:
                    warnings.append(f"Row {idx}: parent '{pval}' not found")
        if parent_temp is None and parent_id is None:
            parent_temp, parent_id = (feature_temp, feature_id) if (feature_temp or feature_id) else (epic_temp, epic_id)

        type_key = _fuzzy_key(get(row, "type"), types, default=leaf) if mapping.get("type") else leaf
        if parent_temp or parent_id:
            parent_level = level_of.get(
                (containers.get(parent_temp) or {}).get("type") if parent_temp else db.get(WorkItem, parent_id).type,
                -1,
            )
            if level_of.get(type_key, 99) <= parent_level:
                lower = [k for k, lv in level_of.items() if lv > parent_level]
                type_key = min(lower, key=lambda k: level_of[k]) if lower else leaf
        spec: dict[str, Any] = {
            "temp_id": f"row:{idx}",
            "type": type_key,
            "title": str(title).strip()[:300],
            "description": (str(get(row, "description")).strip() if get(row, "description") not in (None, "") else None),
            "acceptance_criteria": (
                str(get(row, "acceptance_criteria")).strip() if get(row, "acceptance_criteria") not in (None, "") else None
            ),
            "status": _fuzzy_key(get(row, "status"), workflow, default=default_status, extra=status_extra),
            "priority": _fuzzy_key(get(row, "priority"), _PRIORITY_OPTS, default="medium", extra=_PRIORITY_ALIASES),
            "points": _parse_points(get(row, "points")),
            "labels": _split_labels(get(row, "labels")),
            "due_date": (d.isoformat() if (d := _parse_date(get(row, "due_date"))) else None),
            "source": "excel",
            "source_ref": f"{file_token}:{idx}",
            "_ext_key": _norm(get(row, "key")) or None,
        }
        if mapping.get("assignee"):
            spec["assignee_id"] = resolve_assignee(get(row, "assignee"))
        if parent_temp:
            spec["parent_temp_id"] = parent_temp
        elif parent_id:
            spec["parent_id"] = parent_id
        if get(row, "key") not in (None, ""):
            spec["custom"] = {"import_key": str(get(row, "key"))}
        specs.append(spec)

    ordered = list(containers.values()) + specs
    for s in ordered:
        s.pop("_ext_key", None)
    cleaned = [{k: v for k, v in s.items() if v is not None} for s in ordered]
    created = c.bulk_create(db, project, cleaned, user) if cleaned else []
    db.flush()
    return {"created": created, "skipped": skipped, "warnings": warnings, "sheet": chosen}
